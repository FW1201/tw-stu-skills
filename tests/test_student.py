import importlib.util,json,subprocess,sys
from pathlib import Path
from datetime import datetime
import pytest
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
def module(skill,file):
 p=ROOT/('tw-stu-'+skill)/'scripts';sys.path.insert(0,str(p));spec=importlib.util.spec_from_file_location(skill.replace('-','_')+'_'+file,p/file);m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m);return m
@pytest.mark.parametrize('kind',['course_result','diverse','autobiography'])
def test_portfolio_content_consumed_each_mode(tmp_path,kind):
 f=module('learning-portfolio','generate_portfolio.py');data={'mode':'content','type':kind,'title':'我的作品','sections':[{'heading':'真實反思','text':'我修改兩次圖表，保留三筆訪談比較。','evidence_ids':['E1']}],'evidence':[{'id':'E1','locator':'作品第2頁','excerpt':'三筆比較'}]};out=tmp_path/(kind+'.docx');f.generate(data,out);doc=Document(out);assert '我修改兩次圖表' in '\n'.join(p.text for p in doc.paragraphs);assert len(doc.tables)==1
 data['sections'][0]['text']='我只修改一次圖表。';f.generate(data,tmp_path/'changed.docx');text='\n'.join(p.text for p in Document(tmp_path/'changed.docx').paragraphs);assert '我只修改一次圖表' in text and '修改兩次' not in text
@pytest.mark.parametrize('kind',['course_result','diverse','autobiography'])
def test_framework_explicit_and_nonempty(tmp_path,kind):
 script=ROOT/'tw-stu-learning-portfolio/scripts/generate_portfolio.py';out=tmp_path/'x.docx';r=subprocess.run([sys.executable,str(script),'--framework','--type',kind,'--output',str(out)],cwd=tmp_path,capture_output=True);assert r.returncode==0 and len(Document(out).paragraphs)>=10

def test_portfolio_no_implicit_framework(tmp_path):
 out=tmp_path/'x.docx';r=subprocess.run([sys.executable,str(ROOT/'tw-stu-learning-portfolio/scripts/generate_portfolio.py'),'--output',str(out)],capture_output=True);assert r.returncode==2 and not out.exists()

def schedule():return {'plan_id':'P1','timezone':'Asia/Taipei','generated_at':'2026-10-02T12:00:00+08:00','buffer_fraction':0,'chunk_minutes':20,'slots':[{'start':'2026-10-03T18:00:00+08:00','end':'2026-10-03T19:00:00+08:00'}],'blocked':[{'start':'2026-10-03T18:20:00+08:00','end':'2026-10-03T18:40:00+08:00'}],'tasks':[{'id':'M','title':'數學主動回憶','minutes':30},{'id':'E','title':'英文','minutes':30}]}
def test_schedule_capacity_blocked_and_stable_uids():
 f=module('study-planner','plan.py').plan;d=schedule();r=f(d);assert r['capacity_minutes']==40 and r['scheduled_minutes']==40 and sum(x['minutes'] for x in r['unscheduled'])==20
 assert len({e['uid'] for e in r['events']})==len(r['events']);assert r==f(d)
 for e in r['events']:assert e['end']<='2026-10-03T18:20:00+08:00' or e['start']>='2026-10-03T18:40:00+08:00'
 assert 'BEGIN:VCALENDAR' in r['ics'] and 'DTSTART:20261003T100000Z' in r['ics']

def test_schedule_overlap_and_deadline():
 f=module('study-planner','plan.py').plan;d=schedule();d['slots'].append(d['slots'][0])
 with pytest.raises(ValueError):f(d)
 d=schedule();d['tasks'][0]['deadline']='2026-10-02T19:00:00+08:00';r=f(d);assert next(x for x in r['unscheduled'] if x['task_id']=='M')['minutes']==30

def test_ics_utf8_folding():
 f=module('study-planner','plan.py').fold;s=f('SUMMARY:'+'中文'*60);assert all(len(line.encode())<=75 for line in s.split('\r\n'));assert s.replace('\r\n ','')=='SUMMARY:'+'中文'*60

def test_concept_input_escaping_and_model():
 f=module('concept-viz','generate_concept.py').render;d={'title':'一次函數','nodes':[{'id':'x','label':'x','description':'<script>不應執行</script>'}],'edges':[],'linear':{'a':2,'b':1,'min':0,'max':5,'step':1},'questions':[{'prompt':'x=2時y?','answer':'5'}]};r=f(d);assert '&lt;script&gt;' in r and '\\u003cscript' in r and 'd.linear.a*Number(x.value)+d.linear.b' in r
 d['edges']=[{'from':'x','to':'missing','label':'錯'}]
 with pytest.raises(ValueError):f(d)
