#!/usr/bin/env python3
"""Real student content, or an explicitly requested blank framework."""
import argparse,json,hashlib,os,tempfile
from pathlib import Path
from docx import Document
from docx.shared import Pt,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
TYPES={'course_result':['課程與目標','學習過程','成果與證據','修訂與反思'],'diverse':['活動與角色','實際投入','作品與證據','反思與下一步'],'autobiography':['自我介紹','真實經歷','學習與探索','選擇與方向']}

def validate(d):
    if d['type'] not in TYPES:raise ValueError('unknown portfolio type')
    if d.get('mode') not in ['content','framework','example']:raise ValueError('explicit mode required')
    if not isinstance(d['title'],str) or not d['title'].strip():raise ValueError('title required')
    if not d['sections']:raise ValueError('nonempty sections required')
    for s in d['sections']:
        if not isinstance(s['heading'],str) or not s['heading'].strip() or not isinstance(s['text'],str):raise ValueError('section heading/text required')
        if d['mode']=='content' and not s['text'].strip():raise ValueError('content sections cannot be blank')
    ids=[e['id'] for e in d.get('evidence',[])]
    if len(ids)!=len(set(ids)):raise ValueError('duplicate evidence ids')
    for s in d['sections']:
        if not set(s.get('evidence_ids',[])).issubset(ids):raise ValueError('unknown evidence reference')
    for e in d.get('evidence',[]):
        if not isinstance(e.get('locator'),str) or not e['locator'].strip():raise ValueError('evidence locator required')

def generate(d,path):
    validate(d);doc=Document();doc.sections[0].top_margin=Cm(2);doc.sections[0].bottom_margin=Cm(2)
    style=doc.styles['Normal'];style.font.name='Microsoft JhengHei';style.font.size=Pt(11);style.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei')
    doc.add_heading(d['title'],0);doc.add_paragraph({'content':'依學生提供內容排版','framework':'空白引導框架，非學生已完成經驗','example':'範例，非正式學習紀錄'}[d['mode']])
    if d.get('grade'):doc.add_paragraph('年段：'+d['grade'])
    for s in d['sections']:
        doc.add_heading(s['heading'],1);doc.add_paragraph(s['text'] if s['text'] else '請以自己的真實經歷與作品填寫。')
        if s.get('evidence_ids'):doc.add_paragraph('證據：'+', '.join(s['evidence_ids']))
    if d.get('evidence'):
        doc.add_heading('證據對照',1);table=doc.add_table(rows=1,cols=3);table.style='Table Grid'
        for cell,text in zip(table.rows[0].cells,['ID','來源／位置','原文摘錄']):cell.text=text
        for e in d['evidence']:
            for cell,text in zip(table.add_row().cells,[e['id'],e['locator'],e.get('excerpt','')]):cell.text=text
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    fd,tmp=tempfile.mkstemp(suffix='.docx',dir=path.parent);os.close(fd)
    try:doc.save(tmp);os.replace(tmp,path)
    finally:
        if os.path.exists(tmp):os.unlink(tmp)
    return doc

def main():
    p=argparse.ArgumentParser();p.add_argument('--input',type=Path);p.add_argument('--content');p.add_argument('--framework',action='store_true');p.add_argument('--example',action='store_true');p.add_argument('--type',choices=TYPES,default='course_result');p.add_argument('--grade');p.add_argument('--target_dept');p.add_argument('--output',type=Path);p.add_argument('--validate-only',action='store_true');a=p.parse_args()
    try:
        if sum([a.input is not None,a.content is not None,a.framework,a.example])!=1:raise ValueError('choose --input, --content, --framework or --example')
        if a.input:d=json.loads(a.input.read_text())
        else:
            mode='content' if a.content is not None else 'framework' if a.framework else 'example'
            d={'mode':mode,'type':a.type,'title':'學習歷程'+('（範例）' if mode=='example' else ''),'grade':a.grade,'sections':[{'heading':'學生原文','text':a.content}] if a.content is not None else [{'heading':h,'text':'示例文字：本週完成一份草稿，仍需以真實內容替換。' if mode=='example' else ''} for h in TYPES[a.type]],'evidence':[]}
        validate(d)
        if a.validate_only:print('PASS portfolio content');return
        if not a.output:raise ValueError('--output required')
        generate(d,a.output);a.output.with_suffix('.validation.json').write_text(json.dumps({'mode':d['mode'],'type':d['type'],'sections':len(d['sections']),'evidence_count':len(d.get('evidence',[])),'input_sha256':hashlib.sha256(a.input.read_bytes()).hexdigest() if a.input else None,'checks_run':['content consumed','evidence references'],'limitations':['排版不代表學校規範或 Office 視覺已驗收。']},ensure_ascii=False,indent=2));print(a.output)
    except (ValueError,KeyError,TypeError,OSError) as e:p.exit(2,str(e)+'\n')
if __name__=='__main__':main()
